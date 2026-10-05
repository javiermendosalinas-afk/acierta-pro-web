import re
import os, sys, json, importlib.util, tempfile, shutil, io, contextlib
RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
spec = importlib.util.spec_from_file_location("S", os.path.join(RAIZ, "herramientas", "inventario_sync.py"))
S = importlib.util.module_from_spec(spec); spec.loader.exec_module(S)

fallas = 0
def ok(c, m):
    global fallas
    if not c: fallas += 1
    print(('OK   ' if c else 'FALLA') + ' ' + m)

print("── PARSEO DE TARJETAS")
html = '''
<div class="property-listing" data-lat="20.6" data-long="-103.3"
 data-popover-data='{"price": "$1,500,000 MXN", "title": "Casa bonita", "location": "Casa en condominio en Providencia, San Pedro Tlaquepaque", "bedrooms": "3", "bathrooms": "2.5", "size": "150 m2", "url": "/property/casa-bonita", "image_url": "https://assets.easybroker.com/property_images/1/2/EB-WX1234.jpg?height=300&version=1&width=450"}'>
 <div class="property-photo"><a><img alt="EB-WX1234" src="x"></a></div>
 <div class="property-info"><p class="property-type">Casa en condominio</p></div></div>
<div class="property-listing" data-lat="20.7" data-long="-103.4"
 data-popover-data='{"price": "$85,100,000 USD", "title": "Terreno sin foto", "location": "Terreno en Zapopan, Zapopan", "size": "500 m2", "url": "/property/terreno"}'>
 <div class="property-info"><p class="property-type">Terreno</p></div></div>'''
t = S.parsear_tarjetas(html)
ok(len(t) == 2, "lee las 2 tarjetas")
ok(t[0]["foto"].startswith("https://assets.easybroker.com") and t[0]["codigo_eb"] == "EB-WX1234", "captura la foto y el codigo EB (el bug original)")
ok(t[0]["tipo"] == "casa en condominio" and t[0]["recamaras"] == 3 and t[0]["banos"] == 2.5, "tipo en minusculas, recamaras entero, banos decimal")
ok(t[1]["foto"] is None and t[1]["moneda"] == "USD", "sin foto queda None y detecta USD")

print("── SEGMENTO")
casos = [
 ("oficina", "Oficina en Providencia", "comercial"), ("nave industrial", "x", "comercial"),
 ("terreno comercial", "x", "comercial"), ("departamento", "Depa", "vivienda"),
 ("terreno", "TERRENO EN VENTA IDEAL PARA PROYECTOS DE USOS MIXTOS COMERCIAL", "comercial"),
 ("terreno", "Terreno en Zapopan", "vivienda"), ("terreno", "Terreno habitacional en coto", "vivienda"),
 ("edificio", "Se vende Hotel en Av. Vallarta", "comercial"), ("edificio", "Edificio de departamentos residencial", "vivienda"),
 ("edificio", "Edificio sin pista", "vivienda"), ("casa con uso de suelo", "Casa para oficinas corporativas", "vivienda"),
 ("otro", "algo raro", "vivienda"), ("local en centro comercial", "x", "comercial"),
]
for tipo, tit, esp in casos:
    seg, _ = S.clasificar_segmento(tipo, tit)
    ok(seg == esp, f"{tipo!r} / {tit[:40]!r} -> {seg} (esperado {esp})")

print("── GUARDAS DE PRECIO (casos REALES de hoy)")
def f(op, tipo, precio, m2, moneda=None, seg=None):
    seg = seg or S.clasificar_segmento(tipo, "")[0]
    d = {"operacion": op, "tipo": tipo, "precio": precio, "m2": m2, "segmento": seg}
    if moneda: d["moneda"] = moneda
    return d
reales = [
 (f("VENTA","casa en condominio",31_500_000_000,508), "bloquea", "casa 508m2 a 31,500 millones"),
 (f("VENTA","departamento",1_850_000_000,53), "bloquea", "depto 53m2 a 1,850 millones"),
 (f("VENTA","terreno",507_000_000,39), "bloquea", "terreno 39m2 a 507 millones"),
 (f("VENTA","terreno",388_656_000,4.8), "aviso", "terreno 4.8m2 a 388M: m2 dudoso, se juzga sin m2 (aviso)"),
 (f("VENTA","terreno",432_000_000,520_000), "ok", "terreno 520,000m2 a 432M (831 $/m2, legitimo)"),
 (f("VENTA","terreno",2_500_000_000,1_000_000), "ok", "terreno 1,000,000m2 a 2,500M (2,500 $/m2, legitimo)"),
 (f("VENTA","edificio",450_000_000,12185), "ok", "edificio 12,185m2 a 450M (legitimo)"),
 (f("RENTA","casa",5_000_000,180), "bloquea", "renta casa 180m2 a 5M/mes"),
 (f("RENTA","departamento",4_548_289,53), "bloquea", "renta depto 53m2 a 4.5M/mes"),
 (f("RENTA","nave industrial",3_644_912,23200), "ok", "renta nave 23,200m2 a 3.6M (157 $/m2, legitimo)"),
 (f("RENTA","terreno",1_462_500,32500), "ok", "renta de terreno 32,500m2 a 1.46M/mes (45 $/m2, legitimo; no es renta de vivienda)"),
 (f("RENTA","oficina",25_000,20), "ok", "oficina chica 20m2 a 25,000/mes (1,250 $/m2, plausible)"),
 (f("RENTA","local comercial",300_000,30), "bloquea", "local 30m2 a 300,000/mes (10,000 $/m2/mes: imposible)"),
 (f("RENTA","local comercial",90_000,30), "aviso", "local 30m2 a 90,000/mes (3,000 $/m2/mes: caro, se publica con aviso)"),
 (f("VENTA","edificio",85_100_000_000,12185,"USD"), "bloquea", "USD 85,100 millones"),
 (f("VENTA","edificio",85_000_000,8000,"USD"), "aviso", "USD 85 millones (aviso, no se oculta)"),
 (f("VENTA","departamento",2_500_000,70), "ok", "depto normal 2.5M"),
 (f("VENTA","departamento",5_000,70), "bloquea", "venta de $5,000"),
 (f("VENTA","terreno",90_000,200), "aviso", "terreno barato $90,000 (aviso)"),
 (f("RENTA","departamento",800,50), "bloquea", "renta de $800"),
 (f("RENTA","departamento",2_000,50), "aviso", "renta de $2,000 (aviso)"),
 (f("RENTA","departamento",15_000,70), "ok", "renta normal 15,000"),
 (f("VENTA","casa",None,100), "bloquea", "sin precio"),
]
for fila, esp, desc in reales:
    aviso_m2 = S.sanear_m2(fila)                    # igual que en el proceso real
    nivel, motivo = S.evaluar_precio(fila)
    if aviso_m2 and nivel == "ok":
        nivel, motivo = "aviso", aviso_m2
    ok(nivel == esp, f"{desc} -> {nivel}" + (f" ({motivo})" if motivo else ""))

print("── M2 DUDOSOS")
caso = {"operacion":"VENTA","tipo":"casa","precio":3_295_000,"m2":1.0,"segmento":"vivienda"}
aviso = S.sanear_m2(caso)
ok(aviso and caso["m2"] is None and S.evaluar_precio(caso)[0] == "ok", "casa de 3.3M con '1 m2': se publica SIN m2 (el precio es normal), con aviso")
caso2 = {"operacion":"VENTA","tipo":"casa","precio":3_000_000,"m2":120.0,"segmento":"vivienda"}
ok(S.sanear_m2(caso2) is None and caso2["m2"] == 120.0, "m2 normales no se tocan")

print("── PRECIO POR M2 (casos REALES de la corrida)")
def pm(op, tipo, precio, m2, moneda=None, texto=""):
    d = {"operacion": op, "tipo": tipo, "precio": precio, "m2": m2, "_precio_texto": texto, "segmento": S.clasificar_segmento(tipo, "")[0]}
    if moneda: d["moneda"] = moneda
    return d
casos_pm = [
 (pm("VENTA","terreno",18_000,164.65), "inferido", 2_963_700, "terreno 164 m2 a $18,000 -> $2.96M"),
 (pm("RENTA","oficina",670,218), "inferido", 146_060, "oficina 218 m2 a $670/m2 -> $146,060/mes"),
 (pm("RENTA","bodega industrial",120,10750), "inferido", 1_290_000, "bodega 10,750 m2 a $120 -> $1.29M/mes"),
 (pm("RENTA","nave industrial",6.3,3360,"USD"), "inferido", 21_168, "nave USD 6.3/m2 x 3,360 m2 -> USD 21,168/mes"),
 (pm("VENTA","terreno comercial",4_550,14_417,texto="$4,550 MXN por m²"), "anuncio", 65_597_350, "el anuncio dice 'por m2'"),
]
for fila, via_esp, total_esp, desc in casos_pm:
    via, motivo, bloquea = S.resolver_precio_por_m2(fila)
    ok(via == via_esp and fila["precio"] == total_esp and fila.get("pm2_pub") and not bloquea, f"{desc} -> {via}, total {fila['precio']:,.0f}")
no_conv = [
 (pm("VENTA","casa",15_000,70), "casa de vivienda con precio bajo NO se convierte (es error, no precio por m2)"),
 (pm("VENTA","terreno",18_000,50), "terreno de 50 m2 (< 100) no se infiere: queda para las guardas"),
 (pm("RENTA","departamento",670,60), "departamento a $670 NO se convierte"),
 (pm("VENTA","terreno",2_960_000,164), "precio ya total no se toca"),
 (pm("RENTA","oficina",25_000,20), "oficina 20 m2 a $25,000/mes (total plausible) no se toca"),
]
for fila, desc in no_conv:
    antes = fila["precio"]; via, _m, _b = S.resolver_precio_por_m2(fila)
    ok(via is None and fila["precio"] == antes, desc)
fila = pm("VENTA","terreno",5_000,None,texto="$5,000 MXN por m2")
via, motivo, bloquea = S.resolver_precio_por_m2(fila)
ok(bloquea and "superficie" in motivo, "anuncio por m2 sin superficie: no se puede calcular el total -> se reporta")
conv = pm("VENTA","terreno",18_000,164.65); S.resolver_precio_por_m2(conv)
ok(S.evaluar_precio(conv)[0] == "ok", "tras convertir, el precio total pasa las guardas")

print("── PRECIO POR M2: DOS LECTURAS (el anuncio dice 'por m2' pero a veces el numero ya es el total)")
casos_total = [
 (pm("VENTA","terreno",20_189_800,1009.49,texto="$20,189,800 MXN por m²"), "terreno 1,009 m2 '$20,189,800 por m2' -> es el total"),
 (pm("RENTA","oficina",70_000,148,texto="$70,000 MXN por m²"), "oficina 148 m2 '$70,000 por m2' -> es la renta total"),
 (pm("RENTA","bodega comercial",294_598,1847,texto="$294,598 MXN por m²"), "bodega 1,847 m2 '$294,598 por m2' -> es la renta total"),
 (pm("VENTA","departamento",6_215_000,68,texto="$6,215,000 MXN por m²"), "departamento 68 m2 '$6,215,000 por m2' -> es el total"),
 (pm("RENTA","casa en condominio",9_000,93,texto="$9,000 MXN por m²"), "casa 93 m2 '$9,000 por m2' de renta -> es el total"),
]
for fila, desc in casos_total:
    antes = fila["precio"]; via, motivo, bloquea = S.resolver_precio_por_m2(fila)
    ok(via == "total" and fila["precio"] == antes and "pm2_pub" not in fila and not bloquea, f"{desc} (precio se queda en {fila['precio']:,.0f})")
fila = pm("RENTA","bodega industrial",120,10750,texto="$120 MXN por m²")
via, _m, _b = S.resolver_precio_por_m2(fila)
ok(via == "anuncio" and fila["precio"] == 1_290_000, "bodega '$120 por m2' x 10,750 m2 = $1.29M/mes (si es por m2)")
fila = pm("VENTA","terreno",100,1000,texto="$100 MXN por m²")
via, _m, _b = S.resolver_precio_por_m2(fila)
ok(via is None and fila["precio"] == 100, "si ninguna lectura es creible no se toca (las guardas deciden)")
fila = pm("VENTA","terreno",18_000,164.65,texto="$18,000 MXN por m²")
S.resolver_precio_por_m2(fila)
ok(fila["precio"] == 2_963_700 and fila["pm2_pub"] == 18_000, "terreno '$18,000 por m2' x 164.65 m2 sigue convirtiendose a $2.96M")

print("── ARCHIVO PARA CHATGPT")
base = {"eb": "EB-WX1234", "operacion": "VENTA", "segmento": "vivienda", "tipo": "casa en condominio", "municipio": "Zapopan",
        "colonia": "Valle Real", "titulo": "Casa en Valle Real", "precio": 9_900_000, "m2": 247.0, "recamaras": 3, "banos": 3.5,
        "niveles": 2, "foto": "https://assets.easybroker.com/property_images/1/2/EB-WX1234.jpg?height=300&version=9&width=450",
        "liga": "https://www.aciertamax.com/property/casa-valle-real", "lat": 20.7, "lon": -103.4}
r = S.fila_chatgpt(base, "2026-10-04")
from urllib.parse import unquote, urlparse, parse_qs
texto_wa = parse_qs(urlparse(r["liga_whatsapp"]).query)["text"][0]
ok("EB-WX1234" in texto_wa and "venta" in texto_wa and "ficha.html?eb=EB-WX1234&op=VENTA" in texto_wa and r["liga_whatsapp"].startswith("https://wa.me/523333777337?text="),
   "la liga de WhatsApp trae la clave EB, la operacion y la ficha (asi la lee Wati/MAX)")
ok(r["precio_texto"] == "$9,900,000 MXN" and r["rango_precio"] == "$8 a $12 millones", f"precio redactado y rango: {r['precio_texto']} / {r['rango_precio']}")
ok(r["datos_clave"] == "3 recámaras · 3.5 baños · 247 m² · 2 niveles", f"datos clave: {r['datos_clave']}")
ok(r["imagen_titular"] == "Casa en condominio en venta" and r["imagen_ubicacion"] == "Valle Real, Zapopan" and r["imagen_precio"] == "$9,900,000 MXN", "textos para la imagen")
ok("width=1200" in r["foto_principal"] and "width=450" in r["foto_miniatura"], "foto grande y miniatura de respaldo")
ok(r["hashtags_sugeridos"].startswith("#AciertaMax #Zapopan #CasaEnVenta") and "#ValleReal" in r["hashtags_sugeridos"], f"hashtags: {r['hashtags_sugeridos']}")
ok(r["apto_para_publicar"] == "si" and r["grupo_tipo"] == "casa" and r["precio_por_m2"] == 40081, "apto = si, grupo y precio por m2")
rr = S.fila_chatgpt(dict(base, _revisar=True, _nota="precio dudoso"), "2026-10-04")
ok(rr["apto_para_publicar"] == "revisar" and "dudoso" in rr["nota_calidad"], "una ficha con duda de precio sale como 'revisar'")
rs = S.fila_chatgpt(dict(base, m2=None, tipo="terreno", recamaras=None, banos=None, niveles=None), "2026-10-04")
ok(rs["m2"] == "" and "no mencionar m²" in rs["nota_calidad"].lower() and rs["datos_clave"] == "", "sin m2 confiable: dato vacio y nota para no mencionarlo")
rn = S.fila_chatgpt(dict(base, eb=""), "2026-10-04")
ok(rn["apto_para_publicar"] == "revisar", "sin clave EB no es apta (Wati no podria identificarla)")
rent = S.fila_chatgpt(dict(base, operacion="RENTA", precio=18_500, m2=90.0), "2026-10-04")
ok(rent["precio_texto"] == "$18,500 MXN al mes" and rent["rango_precio"] == "$12,000 a $18,000 al mes".replace("$12,000 a $18,000", "$18,000 a $25,000") and "op=RENTA" in rent["liga_ficha_acierta_pro"], f"renta: {rent['precio_texto']} / {rent['rango_precio']}")
com = S.fila_chatgpt(dict(base, segmento="comercial", tipo="bodega industrial", recamaras=None, banos=None, niveles=None), "2026-10-04")
ok("#InmueblesComerciales" in com["hashtags_sugeridos"] and com["grupo_tipo"] == "bodega", "comercial: hashtags y grupo propios")

import tempfile as _tf, os as _os
_d = _tf.mkdtemp(); _os.chdir(_d)
n = S.escribir_chatgpt([dict(base), dict(base, eb="EB-ZZ0001", segmento="comercial", tipo="oficina", operacion="RENTA", precio=25_000, m2=20.0)], "2026-10-04")
import csv as _csv
filas_csv = list(_csv.DictReader(open("inventario-chatgpt.csv", encoding="utf-8")))
ok(n == 2 and len(filas_csv) == 2 and set(filas_csv[0].keys()) == {c for c, _d2 in S.COLUMNAS_CHATGPT}, "CSV con todas las columnas documentadas")
from openpyxl import load_workbook
wb = load_workbook("inventario-chatgpt.xlsx")
ok(wb.sheetnames == ["Leeme", "Inventario", "Guia de columnas"], f"Excel con hojas: {wb.sheetnames}")
ok(wb["Inventario"].auto_filter.ref is not None and wb["Inventario"].freeze_panes == "B2" and wb["Inventario"].max_row == 3, "Excel con filtros y encabezado fijo")
ok(wb["Guia de columnas"].max_row == len(S.COLUMNAS_CHATGPT) + 1, "la guia describe cada columna")

ok(S.fila_chatgpt(dict(base, tipo="rancho"), "2026-10-04")["hashtags_sugeridos"].startswith("#AciertaMax #Zapopan #RanchoEnVenta"), "un rancho lleva #RanchoEnVenta, no #CasaEnVenta")
ok(S.fila_chatgpt(dict(base, revisar=True, nota="precio dudoso"), "2026-10-04")["apto_para_publicar"] == "revisar", "las banderas guardadas en data.json (revisar/nota) se respetan al regenerar el archivo")
ok(not re.search(r"recamara|\bbanos\b|\bm2\b", " ".join(str(v) for k, v in S.fila_chatgpt(base, "2026-10-04").items() if k in ("datos_clave", "imagen_datos", "nota_calidad", "rango_precio"))), "los textos para publicar llevan acentos y ñ (recámaras, baños, m²)")

print("── MUNICIPIO / FOTO / FILA")
ok(S.normalizar_municipio("San Pedro Tlaquepaque", "tlaquepaque") == "Tlaquepaque", "San Pedro Tlaquepaque -> Tlaquepaque")
ok(S.normalizar_municipio("Tlajomulco de Zúñiga", "tlajomulco-de-zuniga") == "Tlajomulco de Zúñiga", "Tlajomulco de Zúñiga es parte de la zona")
ok(S.normalizar_municipio("Tlajomulco", "zapopan") == "Tlajomulco de Zúñiga", "'Tlajomulco' tambien se reconoce")
ok(S.normalizar_municipio("Juanacatlán", "zapopan") is None, "un municipio fuera de la zona queda fuera")
ok(S.normalizar_municipio("El Salto", "el-salto") == "El Salto", "El Salto se reconoce (corredor industrial)")
_ts = {"municipio_tarjeta": "El Salto", "slug_municipio": "el-salto", "tipo": "casa", "titulo": "Casa en El Salto", "precio": 1e6,
       "recamaras": 2, "banos": 1, "m2": 80, "codigo_eb": "EB-S1", "href": "x", "foto": "f", "lat": 1, "lon": 1, "colonia": "Centro", "operacion": "VENTA"}
ok(S.construir_fila(_ts, {})[0] is None, "El Salto: la vivienda queda fuera")
_ts.update(tipo="bodega industrial", titulo="Bodega en El Salto")
ok((S.construir_fila(_ts, {})[0] or {}).get("segmento") == "comercial", "El Salto: las bodegas entran como comercial")
ok(S.normalizar_municipio("", "tonala") == "Tonalá", "sin municipio en tarjeta usa el de la URL")
u = "https://assets.easybroker.com/property_images/1/2/EB-X.jpg?height=300&version=9&width=450"
ok(S.foto_con_tamano(u, 1200, 800) == "https://assets.easybroker.com/property_images/1/2/EB-X.jpg?height=800&version=9&width=1200", "reescribe width/height conservando version")
prev = {("EB-WX1234","VENTA"): {"niveles": 2, "foto": "https://viejo/foto.jpg"}}
tt = dict(t[0]); tt["operacion"]="VENTA"; tt["slug_municipio"]="tlaquepaque"; tt["foto"]=None
fila, _ = S.construir_fila(tt, prev)
ok(fila["niveles"] == 2 and fila["foto"] == "https://viejo/foto.jpg", "conserva niveles y foto de la corrida anterior cuando el listado no la trae")
tt2 = dict(t[0]); tt2["operacion"]="VENTA"; tt2["slug_municipio"]="zapopan"; tt2["municipio_tarjeta"]="El Salto"
fila2, motivo = S.construir_fila(tt2, {})
ok(fila2 is None and motivo == "fuera de zona", "una tarjeta de vivienda en El Salto se descarta (El Salto es solo comercial)")

print("── TOPE DE 100 PAGINAS DE EASYBROKER (doble pasada)")
class _Resp:
    def __init__(s, text): s.text = text
def _simular_eb(total, por_pag=18, tope=100):
    precios = list(range(total, 0, -1))   # EB-<total> es la mas cara
    def get(ses, url):
        m = re.search(r"page=(\d+)", url); pag = int(m.group(1)) if m else 1
        lista = precios if "price-desc" in url else precios[::-1]
        if pag > tope: return _Resp("[]|")
        trozo = lista[(pag-1)*por_pag: pag*por_pag]
        sig = f"Siguiente page={pag+1}" if pag < tope and pag*por_pag < total else ""
        return _Resp(json.dumps(trozo) + "|" + sig)
    return get
_orig = (S.get_con_reintentos, S.parsear_tarjetas, S.time.sleep)
S.parsear_tarjetas = lambda html: [{"codigo_eb": f"EB-{n}", "precio": n} for n in json.loads(html.split("|")[0] or "[]")]
S.time.sleep = lambda s: None
for total in (2500, 1000, 4000):
    S.AVISOS_TOPE.clear()
    S.get_con_reintentos = _simular_eb(total)
    with contextlib.redirect_stdout(io.StringIO()):
        filas, completo = S.recorrer(None, "VENTA", "venta", "zapopan", 0)
    unicas = {f["codigo_eb"] for f in filas}
    if total <= 3600:
        ok(completo and len(unicas) == total and not S.AVISOS_TOPE,
           f"{total} fichas en EB: se recuperan todas ({len(unicas)}), incluidas las mas baratas")
    else:
        ok(completo and len(unicas) == 3600 and len(S.AVISOS_TOPE) == 1,
           f"{total} fichas en EB: se recuperan 3600 y el resumen avisa que pueden faltar")
S.get_con_reintentos, S.parsear_tarjetas, S.time.sleep = _orig
S.AVISOS_TOPE.clear()

print("── CORRIDA COMPLETA SIMULADA (con frenos)")
SLUG = {"Guadalajara":"guadalajara","Zapopan":"zapopan","Tlaquepaque":"tlaquepaque","Tonalá":"tonala","Tlajomulco de Zúñiga":"tlajomulco-de-zuniga"}
def tarjeta(eb, op, mun, tipo, precio, m2, foto="http://f/x.jpg", titulo="Prop", moneda="MXN", mun_tarj=None):
    return {"href": f"https://www.aciertamax.com/property/{eb}", "precio": precio, "moneda": moneda, "m2": m2, "recamaras": 3,
            "banos": 2.0, "colonia": "Col", "municipio_tarjeta": mun_tarj if mun_tarj is not None else mun, "tipo": tipo, "titulo": titulo,
            "codigo_eb": eb, "foto": foto, "lat": 20.6, "lon": -103.3, "operacion": op, "slug_municipio": SLUG[mun]}
def correr(tarjetas_por_combo, previo, args=()):
    d = tempfile.mkdtemp(); os.chdir(d)
    json.dump(previo, open("data.json","w"))
    S.recorrer = lambda ses, op, sop, smun, pm: (tarjetas_por_combo.get((op, smun), []), True)
    S.foto_desde_detalle = lambda ses, href: None
    S.probar_foto_grande = lambda filas, muestra=12: (True, "simulada")
    S.time.sleep = lambda s: None
    sys.argv = ["x", *args]
    buf = io.StringIO(); code = 0
    with contextlib.redirect_stdout(buf):
        try: S.main()
        except SystemExit as e: code = e.code
    return d, code, buf.getvalue()

previo = [{"municipio":"Zapopan","operacion":"VENTA","precio":3e6,"titulo":"a","tipo":"casa","recamaras":3,"banos":2.0,"m2":100.0,"niveles":2,"eb":"EB-A1","liga":"l","foto":"f","lat":1.0,"lon":2.0,"colonia":"c"},
          {"municipio":"Tlajomulco de Zúñiga","operacion":"VENTA","precio":3e6,"titulo":"t","tipo":"casa","recamaras":3,"banos":2.0,"m2":100.0,"niveles":None,"eb":"EB-T1","liga":"l2","foto":"f","lat":1.0,"lon":2.0,"colonia":"c"}]
combos = {("VENTA","zapopan"): [
    tarjeta("EB-A1","VENTA","Zapopan","casa",3_000_000,100),
    tarjeta("EB-N1","VENTA","Zapopan","oficina",900_000,40),
    tarjeta("EB-B1","VENTA","Zapopan","casa",31_500_000_000,508),
    tarjeta("EB-F1","VENTA","Zapopan","casa",2_000_000,90,foto=None),
    tarjeta("EB-X1","VENTA","Zapopan","casa",2_000_000,90,mun_tarj="El Salto"),
    tarjeta("EB-A1","VENTA","Zapopan","casa",3_000_000,100),
] + [tarjeta(f"EB-OK{i}","VENTA","Zapopan","departamento",2_000_000+i,60) for i in range(40)],
    ("VENTA","tlajomulco-de-zuniga"): [tarjeta("EB-TJ1","VENTA","Tlajomulco de Zúñiga","casa",1_200_000,80)]}
d, code, out = correr(combos, previo)
data = json.load(open(os.path.join(d, "data.json"), encoding="utf-8"))
ebs = {p["eb"] for p in data}
ok(code == 0, f"la corrida termina bien (codigo {code})")
esperado = {"EB-A1", "EB-N1", "EB-TJ1"} | {f"EB-OK{i}" for i in range(40)}
ok(ebs == esperado, f"publica solo lo valido (sin la bloqueada, la sin foto, la fuera de zona ni la duplicada): {len(ebs)} fichas")
ok(next(p for p in data if p["eb"]=="EB-N1")["segmento"] == "comercial" and next(p for p in data if p["eb"]=="EB-A1")["segmento"] == "vivienda", "cada ficha lleva su segmento")
ok(next(p for p in data if p["eb"]=="EB-A1")["niveles"] == 2, "conserva 'niveles' de la corrida anterior")
ok(any("Tlajomulco" in p["municipio"] for p in data) and not any("Salto" in p["municipio"] for p in data), "Tlajomulco se conserva y la vivienda de El Salto no entra")
ok(set(data[0].keys()) >= {"municipio","operacion","precio","titulo","tipo","recamaras","banos","m2","niveles","eb","liga","foto","lat","lon","colonia","segmento"}, "mantiene todas las llaves que usa el sitio")
an = open(os.path.join(d,"reportes","inventario","anomalias_precio.csv"), encoding="utf-8").read()
ok("EB-B1" in an and "BLOQUEADA" in an, "la ficha con precio imposible queda en el reporte de anomalias")
sf = open(os.path.join(d,"reportes","inventario","sin_foto.csv"), encoding="utf-8").read()
ok("EB-F1" in sf, "la ficha sin foto queda en el reporte de sin foto")
ok(os.path.exists(os.path.join(d,"inventario.csv")) and os.path.exists(os.path.join(d,"inventario-meta.json")), "genera inventario.csv e inventario-meta.json")
ok(os.path.exists(os.path.join(d,"inventario-chatgpt.csv")) and os.path.exists(os.path.join(d,"inventario-chatgpt.xlsx")), "la corrida tambien genera el archivo para ChatGPT (CSV y Excel)")
import csv as _c2
cg = list(_c2.DictReader(open(os.path.join(d,"inventario-chatgpt.csv"), encoding="utf-8")))
ok(len(cg) == len(data) and all(r["codigo_eb"].startswith("EB-") for r in cg), f"el archivo de ChatGPT tiene las {len(cg)} fichas publicadas, todas con clave EB")

print("── FRENO DE FOTOS: si faltan demasiadas, algo fallo -> NO se descartan")
muchas_sin_foto = {("VENTA","zapopan"): [tarjeta(f"EB-S{i}","VENTA","Zapopan","casa",2_000_000+i,90,foto=None) for i in range(30)] +
                                          [tarjeta(f"EB-C{i}","VENTA","Zapopan","casa",2_100_000+i,90) for i in range(10)]}
dd, cc, oo = correr(muchas_sin_foto, [])
dat = json.load(open(os.path.join(dd, "data.json"), encoding="utf-8"))
ok(cc == 0 and len(dat) == 40 and "FRENO" in open(os.path.join(dd,"reportes","inventario","resumen.md"), encoding="utf-8").read(), f"75% sin foto: se conservan las 40 fichas y el resumen lo avisa ({len(dat)})")

print("── FRENO: corrida incompleta NO sobreescribe")
previo_grande = [dict(previo[0], eb=f"EB-P{i}", liga=f"l{i}") for i in range(100)]
d2, code2, out2 = correr({("VENTA","zapopan"): [tarjeta("EB-A1","VENTA","Zapopan","casa",3_000_000,100)]}, previo_grande)
data2 = json.load(open(os.path.join(d2, "data.json"), encoding="utf-8"))
ok(code2 == 2 and len(data2) == 100 and "FRENO" in out2, f"con 1 ficha vs 100 antes: se frena (codigo {code2}) y data.json queda intacto ({len(data2)} fichas)")

print("── PRUEBA RAPIDA (--sin-escribir) no toca data.json")
d3, code3, _ = correr(combos, previo, ("--sin-escribir", "--paginas-max", "2"))
data3 = json.load(open(os.path.join(d3, "data.json"), encoding="utf-8"))
ok(code3 == 0 and len(data3) == len(previo), "modo prueba no sobreescribe data.json")

print(f"\nFALLAS: {fallas}")
sys.exit(1 if fallas else 0)
